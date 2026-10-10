"""Generate synthetic files: parser venv + test-only `pip install xlwt==1.3.0`."""

from pathlib import Path
import zipfile

import openpyxl
from pptx import Presentation
import xlwt

ROOT = Path(__file__).parent
TEXT = "Orion 팀의 문서 파싱 검증입니다."

for name, content in {
    "sample.txt": TEXT + "\n",
    "sample.md": "# Orion handbook\n\n" + TEXT + "\n",
    "sample.csv": 'Product,Owner,Notes\nOrion,김하늘,"첫째 줄\n둘째 줄"\n',
    "sample.json": '{"product":{"name":"Orion","owner":"김하늘","notes":"문서 파싱 검증"}}\n',
    "sample.xml": '<?xml version="1.0" encoding="UTF-8"?>\n<inventory><product><name>Orion</name><owner>김하늘</owner></product></inventory>\n',
}.items():
    (ROOT / name).write_text(content, encoding="utf-8")


def archive(name, files):
    with zipfile.ZipFile(ROOT / name, "w", zipfile.ZIP_DEFLATED) as target:
        for path, content in files.items():
            info = zipfile.ZipInfo(path, (2026, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_STORED if path == "mimetype" else zipfile.ZIP_DEFLATED
            target.writestr(info, content)


archive("sample.docx", {
    "[Content_Types].xml": '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="xml" ContentType="application/xml"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>',
    "_rels/.rels": '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>',
    "word/document.xml": f'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>Orion handbook</w:t></w:r></w:p><w:p><w:r><w:t>{TEXT}</w:t></w:r></w:p></w:body></w:document>',
    "word/styles.xml": '<w:styles xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:style w:type="paragraph" w:styleId="Heading1"><w:name w:val="heading 1"/></w:style></w:styles>',
    "word/_rels/document.xml.rels": '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rIdStyles" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/></Relationships>',
})

presentation = Presentation()
slide = presentation.slides.add_slide(presentation.slide_layouts[1])
slide.shapes.title.text = "Orion handbook"
slide.placeholders[1].text = TEXT
slide.notes_slide.notes_text_frame.text = "Review the source document."
presentation.save(ROOT / "sample.pptx")

workbook = openpyxl.Workbook()
workbook.active.title = "Orion inventory"
workbook.active.append(["Product", "Owner"])
workbook.active.append(["Orion", "김하늘"])
workbook.create_sheet("Second sheet").append(["Polaris", "박별"])
workbook.save(ROOT / "sample.xlsx")

long_table = openpyxl.Workbook()
long_table.active.title = "Inventory"
long_table.active.append(["Product", "Owner"])
for index in range(180):
    long_table.active.append([f"Orion-{index}", f"Owner-{index}"])
long_table.save(ROOT / "table.xlsx")

legacy = xlwt.Workbook()
sheet = legacy.add_sheet("Orion inventory")
for row, values in enumerate([["Product", "Owner"], ["Orion", "김하늘"]]):
    for col, value in enumerate(values):
        sheet.write(row, col, value)
legacy.save(str(ROOT / "sample.xls"))

archive("sample.epub", {
    "mimetype": "application/epub+zip",
    "META-INF/container.xml": '<container xmlns="urn:oasis:names:tc:opendocument:xmlns:container" version="1.0"><rootfiles><rootfile full-path="book/content.opf" media-type="application/oebps-package+xml"/></rootfiles></container>',
    "book/content.opf": '<package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="book"><metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:title>Orion handbook</dc:title><dc:language>ko</dc:language><dc:identifier id="book">orion-fixture</dc:identifier><meta property="dcterms:modified">2026-01-01T00:00:00Z</meta></metadata><manifest><item id="chapter" href="chapter.xhtml" media-type="application/xhtml+xml"/><item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/></manifest><spine><itemref idref="chapter"/></spine></package>',
    "book/chapter.xhtml": f'<html xmlns="http://www.w3.org/1999/xhtml" lang="ko" xml:lang="ko"><head><title>Orion handbook</title></head><body><h1>Orion handbook</h1><p>{TEXT}</p></body></html>',
    "book/nav.xhtml": '<html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops" lang="ko" xml:lang="ko"><head><title>목차</title></head><body><nav epub:type="toc"><h1>목차</h1><ol><li><a href="chapter.xhtml">Orion handbook</a></li></ol></nav></body></html>',
})

(ROOT / "sample.html").write_text(f'<html><head><title>Orion</title><style>hidden-style</style></head><body><h1>Orion handbook</h1><p>{TEXT}</p><table><tr><th>Product</th><th>Owner</th></tr><tr><td>Orion</td><td>김하늘</td></tr></table><script>hidden-script</script></body></html>')


def pdf(name, text):
    stream = f"BT /F1 18 Tf 50 750 Td ({text}) Tj ET".encode()
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>", b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        f"<< /Length {len(stream)} >>\nstream\n".encode() + stream + b"\nendstream"]
    output = b"%PDF-1.4\n"
    offsets = [0]
    for index, obj in enumerate(objects, 1):
        offsets.append(len(output))
        output += f"{index} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref = len(output)
    output += b"xref\n0 6\n0000000000 65535 f \n"
    output += b"".join(f"{offset:010} 00000 n \n".encode() for offset in offsets[1:])
    output += f"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    (ROOT / name).write_bytes(output)


pdf("sample.pdf", "Orion document conversion fixture")
pdf("empty.pdf", "")
archive("wrong.docx", {"plain.txt": "not a Word document"})
archive("oversized.docx", {"word/document.xml": "x" * (64 * 1024 * 1024 + 1)})
